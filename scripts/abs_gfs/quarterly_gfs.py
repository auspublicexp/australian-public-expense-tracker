"""Build an isolated, quarterly-derived FY summary; never replace annual GFS.

Download with fetch_quarterly_gfs.py, then run this script from any directory.
All output paths are explicit/local. No deployment or cloud mirroring occurs.
"""
from argparse import ArgumentParser
from pathlib import Path
import csv
import hashlib
import json
import math
import sys
from html.parser import HTMLParser
from html import escape

import openpyxl
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'scripts'))
from chart_helpers import add_brand_logo
from validation_report import write_validation_report

RELEASE_URL = 'https://www.abs.gov.au/statistics/economy/government/government-finance-statistics-australia/jun-2026'
ARTICLE_URL = 'https://www.abs.gov.au/articles/insights-government-finance-statistics-june-2026'
MEASURES = {'revenue': 'Total GFS revenue', 'expenses': 'Total GFS expenses',
            'net_operating_balance': 'GFS NET OPERATING BALANCE', 'taxation_revenue': 'Taxation revenue'}
LABELS = {'revenue': 'Revenue', 'expenses': 'Expenses', 'net_operating_balance': 'Net operating balance',
          'taxation_revenue': 'Taxation revenue'}

class Tables(HTMLParser):
    """Read the published accessible HTML tables without optional HTML engines."""
    def __init__(self):
        super().__init__()
        self.tables, self.table, self.row, self.cell = [], None, None, None
    def handle_starttag(self, tag, attrs):
        if tag == 'table': self.table = []
        elif tag == 'tr' and self.table is not None: self.row = []
        elif tag in ('td', 'th') and self.row is not None: self.cell = []
    def handle_data(self, data):
        if self.cell is not None: self.cell.append(data)
    def handle_endtag(self, tag):
        if tag in ('td', 'th') and self.cell is not None:
            self.row.append(' '.join(''.join(self.cell).split())); self.cell = None
        elif tag == 'tr' and self.row is not None:
            self.table.append(self.row); self.row = None
        elif tag == 'table' and self.table is not None:
            self.tables.append(self.table); self.table = None

def row_values(sheet, label, quarters):
    headers = {}
    for cell in sheet[6]:
        if cell.value is not None:
            if cell.value in headers: raise ValueError('Duplicate quarter header')
            headers[cell.value] = cell.column
    matches = [r for r in range(8, sheet.max_row + 1)
               if str(sheet.cell(r, 1).value).strip().casefold() == label.casefold()]
    if len(matches) != 1: raise ValueError(f'{sheet.title}: expected exactly one {label}')
    result = []
    for quarter in quarters:
        if quarter not in headers: raise ValueError(f'Missing quarter: {quarter}')
        col = headers[quarter]
        if sheet.cell(7, col).value != '$m': raise ValueError('Expected AUD millions')
        value = sheet.cell(matches[0], col).value
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
            raise ValueError(f'Non-numeric or missing source value: {label}/{quarter}')
        result.append((value, sheet.cell(matches[0], col).coordinate))
    return result

def article_row(tables, heading, fy):
    matches = [t for t in tables if t and heading in t[0]]
    if len(matches) != 1: raise ValueError(f'Unexpected article table: {heading}')
    table = matches[0]
    rows = [r for r in table[1:] if r and r[0] == fy]
    if len(rows) != 1 or len(rows[0]) != len(table[0]): raise ValueError('Missing/duplicate article year')
    return dict(zip(table[0], rows[0]))

def number(value): return float(value.replace(',', ''))

def calculate(workbook, article):
    book = openpyxl.load_workbook(workbook, data_only=True)
    sheet = book['Table 1']
    if sheet['A2'].value != 'Government Finance Statistics, Australia - Jun 2026':
        raise ValueError('This release-specific builder requires the June 2026 source vintage')
    if str(sheet['A4'].value).strip() != 'Table 1. All Levels of Government, Operating Statement - General Government':
        raise ValueError('Unexpected sector/table title')
    quarters = [c.value for c in sheet[6] if c.value is not None]
    if len(quarters) != 4: raise ValueError('Expected four quarters')
    end_year = int(quarters[-1][-4:])
    expected = [f'Sep Qtr {end_year-1}', f'Dec Qtr {end_year-1}', f'Mar Qtr {end_year}', f'Jun Qtr {end_year}']
    if quarters != expected: raise ValueError('Not one complete July-June financial year')
    fy = f'FY{end_year-1}-{str(end_year)[-2:]}'
    records, totals = [], {}
    for key, label in MEASURES.items():
        values = row_values(sheet, label, quarters)
        totals[key] = sum(v for v, _ in values)
        for q, (value, cell) in zip(quarters, values):
            records.append(dict(financial_year=fy, quarter=q, measure=key, value_millions=value,
                                source_sheet='Table 1', source_cell=cell, source_file=workbook.name))
    previous = [q.replace(str(end_year-1), str(end_year-2)) if q.startswith(('Sep','Dec'))
                else q.replace(str(end_year),str(end_year-1)) for q in quarters]
    fiscal = book['Table 15']
    if fiscal['A4'].value != 'Table 15. All Levels of Government, Key Measures and Fiscal Aggregates' or fiscal['A14'].value != 'Key fiscal aggregates, all levels of general government, original':
        raise ValueError('Unexpected Table 15 classification or series basis')
    current_tax = sum(v for v,_ in row_values(fiscal, 'Taxation revenue', quarters))
    current_nob = sum(v for v,_ in row_values(fiscal, 'Net operating balance', quarters))
    if current_tax != totals['taxation_revenue'] or current_nob != totals['net_operating_balance']:
        raise ValueError('Tables 1 and 15 do not reconcile')
    prior_values = row_values(fiscal, 'Taxation revenue', previous)
    prior_tax = sum(v for v,_ in prior_values)
    for q,(v,c) in zip(previous,prior_values):
        records.append(dict(financial_year=f'FY{end_year-2}-{str(end_year-1)[-2:]}',quarter=q,
                            measure='taxation_revenue',value_millions=v,source_sheet='Table 15',source_cell=c,source_file=workbook.name))
    totals['taxation_increase'] = current_tax-prior_tax
    growth = (current_tax/prior_tax-1)*100
    identity = totals['revenue']-totals['expenses']-totals['net_operating_balance']
    if abs(identity) > 6: raise ValueError('Operating balance identity exceeds published rounding tolerance')
    parser = Tables(); parser.feed(article.read_text(encoding='utf-8'))
    article_fy = f'{end_year-1}/{str(end_year)[-2:]}'
    headline = article_row(parser.tables,'Revenue ($b)',article_fy)
    taxes = article_row(parser.tables,'Personal income tax ($b)',article_fy)
    checks = []
    for key, field in [('revenue','Revenue ($b)'),('expenses','Expenses ($b)'),('net_operating_balance','Net operating balance ($b)')]:
        published = number(headline[field])
        if key == 'expenses': published = abs(published) # ABS plots expenses below zero
        actual = totals[key]/1000
        checks.append(dict(measure=key,calculated_billions=actual,published_billions=published,
                           status='PASS' if abs(actual-published) <= .05 else 'FAIL',method='Latest-vintage quarterly sum versus ABS article'))
    # Targets are validation assertions only; they never replace calculated values.
    for key,actual,expected_value in [('taxation_increase',totals['taxation_increase']/1000,58.6),('taxation_growth_percent',growth,7.0)]:
        checks.append(dict(measure=key,calculated=actual,published=expected_value,
                           status='PASS' if abs(actual-expected_value) <= .05 else 'FAIL',method='Table 15 latest-vintage comparison; article narrative target'))
    tax_items=[]
    for field,target in [('Personal income tax ($b)',344.2),('Company income tax ($b)',151.4),('Goods and services tax (GST) ($b)',98.0)]:
        value=number(taxes[field])
        tax_items.append(dict(measure=field.removesuffix(' ($b)'),value_billions=value,source_url=ARTICLE_URL,
                              basis='ABS published quarterly-derived FY total; not independently decomposable from this workbook'))
        checks.append(dict(measure=field,published_billions=value,requested_billions=target,
                           status='PASS' if value==target else 'FAIL',method='Direct extraction from ABS article; not independent quarterly validation'))
    book.close()
    if any(c['status']!='PASS' for c in checks): raise ValueError(json.dumps(checks,indent=2))
    return dict(financial_year=fy,release='June 2026',release_url=RELEASE_URL,article_url=ARTICLE_URL,
                basis='Sum of four quarterly GFS estimates from the June 2026 release; not the ABS annual release',
                sector='All Australia general government',unit='AUD millions',price_basis='current prices',
                series_type='original',accounting_basis='accrual',quarters=quarters,totals_millions=totals,
                previous_taxation_millions=prior_tax,taxation_growth_percent=growth,tax_items=tax_items,
                rounding_identity_difference_millions=identity,checks=checks,
                source_sha256={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in [workbook,article]}),records

def write_csv(path, rows):
    with path.open('w',encoding='utf-8',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=list(rows[0])); writer.writeheader(); writer.writerows(rows)

def page_content(result):
    rows=''.join(f'<tr><th scope="row">{escape(LABELS[k])}</th><td>${v/1000:,.1f} billion</td></tr>'
                 for k,v in result['totals_millions'].items() if k in LABELS)
    charts=''.join(f'''<div class="card shadow-sm mb-4"><img class="card-img-top" src="/charts/abs_gfs/FY2025-26_quarterly_derived/{name}.svg" alt="{label}"><div class="card-body border-top py-3 d-flex flex-wrap gap-2" role="group" aria-label="Chart downloads"><a class="btn btn-sm btn-outline-secondary" href="/charts/abs_gfs/FY2025-26_quarterly_derived/{name}.png" download>Download PNG</a><a class="btn btn-sm btn-outline-secondary" href="/charts/abs_gfs/FY2025-26_quarterly_derived/{name}.svg" download>Download SVG</a><a class="btn btn-sm btn-outline-secondary" href="/charts/abs_gfs/FY2025-26_quarterly_derived/{name}_data.csv" download>Download CSV</a></div></div>'''
    for name,label in [('01_fiscal_summary','FY2025–26 revenue, expenses and net operating balance'),('02_selected_taxes','FY2025–26 personal income tax, company income tax and GST')])
    return f'''<main class="container py-4 py-md-5">
<a href="/abs-gfs/">Back to ABS Government Finance Statistics</a>
<h1 class="display-5 fw-bold mt-4">FY2025–26 government finances</h1>
<p class="lead">July 2025 to June 2026 — quarterly-derived financial-year results</p>
<div class="alert alert-light border" role="note"><h2 class="h5">Basis and limitations</h2>
<p>These figures sum four quarters from the latest June 2026 GFS release, including revisions to earlier quarters. They cover consolidated All Australia general government: current prices, original series, accrual accounting.</p>
<p>This is not the ABS annual GFS release, which is scheduled for 20 April 2027. Quarterly-derived results may differ from that release. The existing annual expense series and topic explorer remain available through FY2024–25.</p>
<p class="mb-0">Expense-by-purpose data and a separate local-government breakdown are not available in this quarterly workbook. No estimated breakdowns have been added.</p></div>
<table class="table table-striped"><caption>Sum of four quarters; AUD billions, rounded to one decimal place</caption><thead><tr><th scope="col">Measure</th><th scope="col">FY2025–26</th></tr></thead><tbody>{rows}</tbody></table>
<p>Taxation revenue increased {result['taxation_growth_percent']:.1f}% (${result['totals_millions']['taxation_increase']/1000:.1f} billion) compared with FY2024–25, using the revised quarterly values for both years from the June 2026 workbook.</p>
{charts}
<p>Selected tax figures are taken directly from the ABS insights article. The workbook does not report those tax types separately, so APET cannot independently sum their quarters from this source. Values shown are rounded and may not add exactly.</p>
<p><a href="{RELEASE_URL}" target="_blank" rel="noopener noreferrer">Official June 2026 release and workbook</a> · <a href="{ARTICLE_URL}" target="_blank" rel="noopener noreferrer">ABS insights and financial-year figures</a></p>
<p><a href="/charts/abs_gfs/FY2025-26_quarterly_derived/quarterly_values.csv" download>Download underlying quarterly values and source cells</a> · <a href="/charts/abs_gfs/FY2025-26_quarterly_derived/summary.json" download>Download summary, provenance and validation</a></p>
</main>'''

def build(raw, out):
    result, records = calculate(raw/'quarterly.xlsx', raw/'insights.html')
    out.mkdir(parents=True,exist_ok=True)
    (out/'summary.json').write_text(json.dumps(result,indent=2)+'\n',encoding='utf-8')
    (out/'summary.html').write_text(page_content(result),encoding='utf-8')
    write_csv(out/'quarterly_values.csv',records)
    write_csv(out/'01_fiscal_summary_data.csv',[dict(measure=LABELS[k],value_millions=v,basis=result['basis']) for k,v in result['totals_millions'].items() if k in LABELS])
    write_csv(out/'02_selected_taxes_data.csv',result['tax_items'])
    plt.rcParams['svg.fonttype']='none'
    plt.rcParams['svg.hashsalt']='apet-gfs-jun-2026'
    for name,title,labels,values,colors,note in [
        ('01_fiscal_summary','FY2025–26 general government finances',
         [LABELS[k] for k in ['revenue','expenses','net_operating_balance']],
         [result['totals_millions'][k]/1000 for k in ['revenue','expenses','net_operating_balance']],
         ['#1f77b4','#e69138','#777777'],'Latest June 2026 workbook: sum of four quarters; subject to revision'),
        ('02_selected_taxes','FY2025–26 selected taxation revenue',
         [r['measure'].replace('Goods and services tax (GST)','GST') for r in result['tax_items']],
         [r['value_billions'] for r in result['tax_items']],['#1f77b4']*3,
         'ABS-published FY totals from its insights article; selected taxes are not total taxation')]:
        fig,ax=plt.subplots(figsize=(12,7))
        bars=ax.barh(labels,values,color=colors)
        ax.invert_yaxis(); ax.axvline(0,color='#777777',linewidth=.8)
        ax.set_xlim(min(0,min(values))*2,max(values)*1.22)
        ax.bar_label(bars,labels=[f'${v:,.1f}b' if v>=0 else '' for v in values],padding=6,fontsize=12)
        for i,v in enumerate(values):
            if v<0: ax.text(max(values)*.025,i,f'−${abs(v):,.1f}b',va='center',fontsize=12)
        ax.set_xlabel('AUD billions'); ax.set_title(title,loc='left',fontsize=20,pad=35)
        ax.text(0,1.03,'All Australia general government • current prices • original series • accrual',transform=ax.transAxes,fontsize=10,color='#555555')
        ax.spines[['top','right']].set_visible(False)
        add_brand_logo(fig)
        fig.text(.17,.075,'Source: ABS, Government Finance Statistics, June 2026',fontsize=9,url=ARTICLE_URL if name=='02_selected_taxes' else RELEASE_URL)
        fig.text(.17,.05,'auspublicexp.org  |  Australian Public Expense Tracker  |  @auspublicexp',fontsize=9,color='#0d6efd',url='https://auspublicexp.org/')
        fig.text(.17,.025,note,fontsize=8,color='#555555')
        fig.subplots_adjust(left=.23,right=.96,top=.82,bottom=.2)
        for ext in ['png','svg']:
            fig.savefig(out/f'{name}.{ext}',dpi=150,metadata={'Date':None} if ext=='svg' else None)
        plt.close(fig)
    write_validation_report('abs_gfs_quarterly',result['totals_millions'],
        checks=[f"{c['status']}: {c['measure']} — {c['method']}" for c in result['checks']],
        output_dir=out.parent.parent/'validation')
    print(json.dumps({'financial_year':result['financial_year'],'totals_millions':result['totals_millions'],
                      'taxation_growth_percent':result['taxation_growth_percent'],'checks':result['checks']},indent=2))

def main():
    p=ArgumentParser(description=__doc__)
    p.add_argument('--raw-dir',type=Path,default=ROOT/'data/abs_gfs/raw/quarterly/jun-2026')
    p.add_argument('--output-dir',type=Path,default=ROOT/'output/abs_gfs/FY2025-26_quarterly_derived')
    args=p.parse_args(); build(args.raw_dir,args.output_dir)

if __name__=='__main__': main()
